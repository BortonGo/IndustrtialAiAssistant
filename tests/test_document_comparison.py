"""Deterministic revision diff checks and optional large-document benchmark."""
import copy
from pathlib import Path
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'storage_tools'))
from document_comparison import compare_documents


def document(lines=(), blocks=None, identifier='old'):
    return dict(id=identifier, sourcePath=identifier + '.docx', text='\n'.join(lines),
                blocks=blocks or [], warnings=[], pendingImages=0)


class ComparisonTests(unittest.TestCase):
    def test_actual_docx_extraction_then_table_comparison(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest('Run with document_tools/.venv for the DOCX integration check')
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'document_tools'))
        from read_docx import extract_document
        output = Path(__file__).resolve().parents[1] / 'build-qtcreator/test-artifacts'
        output.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=output) as folder:
            extracted = []
            for version, pressure in enumerate(('7.5 bar', '8.0 bar')):
                word = Document()
                word.add_paragraph('5.4 Протокол')
                table = word.add_table(rows=2, cols=2)
                for cells, values in zip(table.rows, [('Поле', 'Значение'), ('Давление', pressure)]):
                    for cell, value in zip(cells.cells, values): cell.text = value
                word.add_paragraph('Конец документа')
                path = Path(folder) / f'revision-{version}.docx'; word.save(path)
                result = extract_document(str(path))
                extracted.append(dict(document(identifier=str(version)), **result))
            report = compare_documents(*extracted)
            self.assertEqual(report['counts'], dict(added=0, removed=0, changed=1))
            self.assertEqual(report['changes'][0]['before']['pageFirst'], 0)
            self.assertIn('7.5 bar', report['changes'][0]['before']['text'])
            self.assertIn('8.0 bar', report['changes'][0]['after']['text'])

    def test_identical_ignores_spacing_and_page_changes(self):
        a = document(blocks=[dict(type='text', text='Насос  PUMP-417: 7.5 bar', pageNumber=1)])
        b = copy.deepcopy(a)
        b['blocks'][0].update(text=' Насос\tPUMP-417: 7.5 bar  ', pageNumber=4)
        self.assertEqual(compare_documents(a, b)['changes'], [])

    def test_values_case_sign_and_leading_zero_are_significant(self):
        for old, new in [('7.5 bar', '8.0 bar'), ('PUMP-417', 'PUMP-418'), ('0x01', '0x1'),
                         ('10 ms', '10 s'), ('-2', '+2'), ('ILS', 'ils')]:
            with self.subTest(old=old):
                report = compare_documents(document([old]), document([new], identifier='new'))
                self.assertEqual(report['counts'], dict(added=0, removed=0, changed=1))
                self.assertEqual(report['changes'][0]['before']['text'], old)
                self.assertEqual(report['changes'][0]['after']['text'], new)

    def test_table_insertion_keeps_duplicates_and_cell_boundaries(self):
        rows = [['Код', 'Значение'], ['01', '7.5'], ['01', '7.5'], ['02', '0']]
        a = document(blocks=[dict(type='table', tableRows=rows, pageNumber=8, tableIndex=3)])
        b = copy.deepcopy(a)
        b['blocks'][0]['tableRows'].insert(2, ['03', '8.0'])
        report = compare_documents(a, b)
        self.assertEqual(report['counts'], dict(added=1, removed=0, changed=0))
        ref = report['changes'][0]['after']
        self.assertEqual(ref['pageFirst'], 8)
        self.assertEqual(ref['tableIndices'], [3])
        self.assertIn('Строка 3', ref['section'])
        self.assertIn('Значение', ref['text'])
        b['blocks'][0]['tableRows'] = rows[:-1]  # removes one row only
        self.assertEqual(compare_documents(a, b)['counts']['removed'], 1)
        b['blocks'][0]['tableRows'] = rows + [['', '']]
        self.assertEqual(compare_documents(a, b)['counts']['added'], 1)

    def test_same_flattened_table_text_does_not_hide_cell_changes(self):
        a = document(blocks=[dict(type='table', tableRows=[['a b', 'c']])])
        b = document(blocks=[dict(type='table', tableRows=[['a', 'b c']])])
        self.assertEqual(compare_documents(a, b)['counts']['changed'], 1)

    def test_named_table_rows_match_with_changed_values_and_insertions(self):
        a = document(blocks=[dict(type='table', tableRows=[['Поле', 'Значение'], ['Давление', '7.5 bar'], ['Скорость', '9600']])])
        b = document(blocks=[dict(type='table', tableRows=[['Поле', 'Значение'], ['Давление', '8.0 bar'], ['CRC', '16'], ['Скорость', '19200']])])
        self.assertEqual(compare_documents(a, b)['counts'], dict(added=1, removed=0, changed=2))

    def test_location_and_direction_are_from_each_snapshot(self):
        a = document(blocks=[dict(type='text', text='5.4 Протокол\nСкорость: 9600', pageNumber=10)])
        b = document(blocks=[dict(type='text', text='5.4 Протокол\nСкорость: 19200\nCRC: 16', pageNumber=11)], identifier='new')
        report = compare_documents(a, b)
        old = next(v['before'] for v in report['changes'] if v['before'])
        new = next(v['after'] for v in report['changes'] if v['after'])
        self.assertEqual((old['documentId'], old['pageFirst']), ('old', 10))
        self.assertEqual((new['documentId'], new['pageFirst']), ('new', 11))
        self.assertEqual(old['section'], '5.4 Протокол')
        inverse = compare_documents(b, a)
        self.assertEqual(report['counts']['added'], inverse['counts']['removed'])
        self.assertEqual(report['counts']['removed'], inverse['counts']['added'])

    def test_images_compare_hash_and_description_and_explain_limit(self):
        a = document(blocks=[dict(type='image', imagePath='a.png', text='Насос', pageNumber=0)])
        b = copy.deepcopy(a)
        for field, value in [('imagePath', 'b.png'), ('text', 'Насос PUMP-417')]:
            b['blocks'][0][field] = value
            result = compare_documents(a, b)
            self.assertEqual(result['counts']['changed'], 1)
            self.assertTrue(result['warnings'])
            self.assertEqual(result['changes'][0]['before']['imagePath'], 'a.png')

    def test_partial_even_with_no_differences(self):
        a = document(['Текст готов'])
        b = copy.deepcopy(a); b['pendingImages'] = 123
        result = compare_documents(a, b)
        self.assertEqual(result['changes'], [])
        self.assertTrue(result['partial'])
        self.assertIn('123', result['warnings'][0])
        b['pendingImages'] = 0; b['warnings'] = ['Страница пропущена']
        self.assertTrue(compare_documents(a, b)['partial'])

    def test_truncation_keeps_totals_and_never_claims_complete_report(self):
        with patch('document_comparison.MAX_CHANGES', 2):
            result = compare_documents(document(['old'] * 5), document(['new'] * 5))
        self.assertEqual(result['counts']['changed'], 5)
        self.assertEqual(len(result['changes']), 2)
        self.assertTrue(result['truncated'])
        self.assertIn('2 изменений из 5', result['warnings'][-1])
        with patch('document_comparison.MAX_REPORT_CHARS', 5):
            result = compare_documents(document(['abc']), document(['def']))
        self.assertTrue(result['truncated'])

    def test_limits_and_empty_documents(self):
        self.assertEqual(compare_documents(document(), document())['changes'], [])
        with patch('document_comparison.MAX_UNITS', 2):
            with self.assertRaises(ValueError): compare_documents(document(['a'] * 3), document())
        with patch('document_comparison.MAX_CHARS', 5):
            with self.assertRaises(ValueError): compare_documents(document(['абв']), document())

    def test_repeated_pages_do_not_turn_one_change_into_thousands(self):
        a = document(['Строка'] * 20000)
        b = document(['Строка'] * 10000 + ['Изменена'] + ['Строка'] * 9999)
        self.assertEqual(compare_documents(a, b)['counts'], dict(added=0, removed=0, changed=1))
        # Force the repetition safeguard with matching interior rows.
        a = document(['start'] + ['x', 'y'] * 2000 + ['end'])
        b = document(['other'] + ['x', 'y'] * 2000 + ['finish'])
        result = compare_documents(a, b)
        self.assertEqual(result['counts']['changed'], 2)
        self.assertTrue(result['warnings'])


def benchmark():
    blocks = [dict(type='text', pageNumber=page + 1,
                   text='\n'.join(f'Страница {page + 1}, устройство PUMP-{row:04d}, значение 7.5 bar'
                                  for row in range(30))) for page in range(500)]
    old = document(blocks=blocks)
    new = copy.deepcopy(old)
    for page in range(0, 500, 5):
        new['blocks'][page]['text'] = new['blocks'][page]['text'].replace('7.5 bar', '8.0 bar', 1)
    started = time.perf_counter()
    report = compare_documents(old, new)
    elapsed = time.perf_counter() - started
    assert report['counts'] == dict(added=0, removed=0, changed=100), report['counts']
    print(f'500 pages / 15000 lines, 100 changes: {elapsed:.3f}s (extracted data, no I/O or OCR)')


if __name__ == '__main__':
    if '--benchmark' in sys.argv:
        benchmark()
    else:
        unittest.main()
