import base64
import io
import json
from pathlib import Path
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, HTTPServer
from types import SimpleNamespace
from unittest.mock import Mock, patch

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "document_tools"))
import read_docx
import read_pdf
import read_document
from document_format import blocks_to_text
from image_processing import ImageAnalyzer, describe_images, save_normalized_image


def make_picture(path):
    image = Image.new("RGB", (900, 360), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 46)
    draw.rectangle((5, 5, 890, 350), outline="red", width=8)
    draw.text((35, 45), "DEVICE: PUMP-417", fill="black", font=font)
    draw.text((35, 125), "PRESSURE: 7.5 bar", fill="black", font=font)
    draw.text((35, 205), "ROOM: 218", fill="black", font=font)
    image.save(path)


def make_pdf(path, picture, rotate=0, native=True):
    """Tiny deterministic PDF fixture without another PDF dependency."""
    stream = io.BytesIO()
    Image.open(picture).convert("RGB").save(stream, "JPEG")
    jpg = stream.getvalue()
    drawing = b"q 270 0 0 108 15 150 cm /Im1 Do Q\n"
    if native:
        drawing += b"BT /F1 12 Tf 15 280 Td (Equipment protocol and pressure readings for test station) Tj ET\n"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 400] /Rotate {rotate} "
         "/Resources << /Font << /F1 4 0 R >> /XObject << /Im1 5 0 R >> >> /Contents 6 0 R >>").encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /XObject /Subtype /Image /Width 900 /Height 360 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /DCTDecode /Length " +
        str(len(jpg)).encode() + b" >>\nstream\n" + jpg + b"\nendstream",
        b"<< /Length " + str(len(drawing)).encode() + b" >>\nstream\n" + drawing + b"endstream"]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    start = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(f"trailer << /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode())
    path.write_bytes(data)


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        (ROOT / "build-document-check").mkdir(exist_ok=True)
        self.tmp = tempfile.TemporaryDirectory(dir=ROOT / "build-document-check")
        self.work = Path(self.tmp.name)
        self.picture = self.work / "plate.png"
        make_picture(self.picture)
        self.images = self.work / "images"

    def tearDown(self):
        self.tmp.cleanup()

    def test_docx_order_cells_merged_nested_header_and_image_only(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        run = paragraph.add_run("Before ")
        run.add_picture(str(self.picture), width=Inches(2))
        run.add_text(" after")
        table = doc.add_table(rows=2, cols=2)
        merged = table.cell(0, 0).merge(table.cell(1, 0))
        merged.text = "Merged label"
        merged.paragraphs[0].add_run().add_picture(str(self.picture), width=Inches(1))
        nested = table.cell(0, 1).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Nested cell"
        nested.cell(0, 0).paragraphs[0].add_run().add_picture(str(self.picture), width=Inches(1))
        doc.sections[0].header.paragraphs[0].text = "Header"
        path = self.work / "cells.docx"
        doc.save(path)
        result = read_docx.extract_document(path, self.images)
        self.assertEqual(result["warnings"], [])
        blocks = result["blocks"]
        self.assertEqual([b["type"] for b in blocks[:3]], ["text", "image", "text"])
        pictures = [b for b in blocks if b["type"] == "image"]
        self.assertEqual(len(pictures), 3)
        self.assertEqual(len(set(b["imagePath"] for b in pictures)), 1)
        self.assertEqual((pictures[1]["tableIndex"], pictures[1]["tableRow"], pictures[1]["tableColumn"]), (0, 0, 0))
        self.assertEqual(pictures[2]["tableIndex"], 1)
        self.assertEqual(result["text"].count("Merged label"), 1)
        self.assertIn("Nested cell", result["text"])
        self.assertIn("Header", result["text"])
        doc = Document()
        doc.add_picture(str(self.picture))
        doc.save(path)
        result = read_docx.extract_document(path, self.images)
        analyzer = Mock()
        analyzer.describe.return_value = ("PUMP-417, 7.5 bar", False)
        describe_images(result["blocks"], analyzer, result["warnings"], lambda text: None)
        self.assertIn("PUMP-417", blocks_to_text(result["blocks"]))

    def test_pdf_photo_scan_rotation_and_blank(self):
        for angle in (0, 90):
            path = self.work / f"photo-{angle}.pdf"
            make_pdf(path, self.picture, angle)
            warnings = []
            blocks = read_pdf.extract_blocks(path, self.images, ROOT / "tools/poppler-26.09.0/Library/bin", warnings, lambda text: None)
            self.assertEqual(warnings, [])
            self.assertIn("Equipment protocol", blocks[0]["text"])
            self.assertEqual(blocks[1]["type"], "image")
            self.assertEqual(blocks[1]["pageNumber"], 1)
            image = Image.open(blocks[1]["imagePath"])
            # Crop must contain the red border of the actual photo, including rotated PDF.
            red = sum(1 for r, g, b in image.get_flattened_data() if r > 180 and g < 90 and b < 90)
            self.assertGreater(red, image.width * 2)
        path = self.work / "scan.pdf"
        Image.open(self.picture).save(path, "PDF")
        warnings = []
        blocks = read_pdf.extract_blocks(path, self.images, ROOT / "tools/poppler-26.09.0/Library/bin", warnings, lambda text: None)
        self.assertEqual([b["type"] for b in blocks], ["image"])
        Image.new("RGB", (300, 400), "white").save(path, "PDF")
        blocks = read_pdf.extract_blocks(path, self.images, ROOT / "tools/poppler-26.09.0/Library/bin", [], lambda text: None)
        self.assertEqual(blocks, [])

    def test_vision_protocol_cache_and_failure(self):
        path = save_normalized_image(self.picture.read_bytes(), self.images)
        requests = []
        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                payload = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                requests.append(payload)
                self.send_response(200)
                self.end_headers()
                self.wfile.write(json.dumps({"choices": [{"message": {"content": "PUMP-417, 7.5 bar"}, "finish_reason": "stop"}]}).encode())
            def log_message(self, *args):
                pass
        server = HTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            analyzer = ImageAnalyzer(f"http://127.0.0.1:{server.server_port}/v1/chat/completions", "local-chat", "fixture", self.work / "cache")
            self.assertEqual(analyzer.describe(path), ("PUMP-417, 7.5 bar", False))
            self.assertEqual(analyzer.describe(path), ("PUMP-417, 7.5 bar", False))
            self.assertEqual(len(requests), 1)
            encoded = requests[0]["messages"][0]["content"][1]["image_url"]["url"].split(",")[1]
            self.assertEqual(base64.b64decode(encoded), Path(path).read_bytes())
        finally:
            server.shutdown()
            server.server_close()
            thread.join()
        blocks = [{"type": "image", "imagePath": path, "pageNumber": 4}]
        warnings = []
        describe_images(blocks, None, warnings, lambda text: None)
        self.assertEqual(blocks[0]["text"], "")
        self.assertIn("страница 4", warnings[0])

    def test_failed_duplicate_images_are_not_requested_repeatedly(self):
        path = save_normalized_image(self.picture.read_bytes(), self.images)
        blocks = [{"type": "image", "imagePath": path, "pageNumber": n} for n in (1, 2)]
        analyzer = Mock()
        analyzer.describe.side_effect = RuntimeError("unsupported image")
        warnings = []
        describe_images(blocks, analyzer, warnings, lambda text: None)
        self.assertEqual(analyzer.describe.call_count, 1)
        self.assertEqual(len(warnings), 2)

    def test_transparent_picture_is_composited_on_white(self):
        source = Image.new("RGBA", (40, 40), (0, 0, 0, 0))
        source.putpixel((20, 20), (255, 0, 0, 255))
        stream = io.BytesIO()
        source.save(stream, "PNG")
        path = save_normalized_image(stream.getvalue(), self.images)
        with Image.open(path) as normalized:
            self.assertEqual(normalized.getpixel((0, 0)), (255, 255, 255))
            self.assertEqual(normalized.getpixel((20, 20)), (255, 0, 0))

    def image_args(self, path):
        return SimpleNamespace(path=path, images_dir=self.images,
                               poppler_dir=self.work / "no-poppler-required",
                               soffice="", vision_url="http://127.0.0.1/unused",
                               vision_profile="fixture")

    def test_standalone_photos_supported_formats_and_exif(self):
        with patch("read_document.ImageAnalyzer") as factory, patch("read_document.progress"):
            factory.return_value.describe.return_value = ("PUMP-417, 7.5 bar", False)
            for extension, format_name in (("PNG", "PNG"), ("JPG", "JPEG"),
                                           ("jpeg", "JPEG"), ("webp", "WEBP"), ("bmp", "BMP")):
                with self.subTest(extension=extension):
                    path = self.work / ("Фото с пробелами." + extension)
                    with Image.open(self.picture) as image:
                        options = {}
                        if format_name == "JPEG":
                            exif = Image.Exif()
                            exif[274] = 6  # A phone photo stored sideways, rotated when displayed.
                            options["exif"] = exif
                        image.save(path, format_name, **options)
                    original = path.read_bytes()
                    result = read_document.extract_document(self.image_args(path))
                    self.assertEqual(result["warnings"], [])
                    self.assertEqual(len(result["blocks"]), 1)
                    block = result["blocks"][0]
                    self.assertEqual((block["type"], block["pageNumber"]), ("image", 0))
                    self.assertIn("PUMP-417", result["text"])
                    self.assertEqual(result["text"], blocks_to_text(result["blocks"]))
                    factory.return_value.describe.assert_called_with(block["imagePath"])
                    with Image.open(block["imagePath"]) as normalized:
                        self.assertEqual(normalized.format, "PNG")
                        self.assertEqual(normalized.mode, "RGB")
                        self.assertEqual(normalized.size, (360, 900) if format_name == "JPEG" else (900, 360))
                    self.assertEqual(path.read_bytes(), original)

    def test_standalone_photo_invalid_animated_and_vision_failure(self):
        with patch("read_document.ImageAnalyzer") as factory, patch("read_document.progress"):
            broken = self.work / "broken.jpg"
            broken.write_bytes(b"not an image")
            with self.assertRaisesRegex(ValueError, "Не удалось прочитать изображение"):
                read_document.extract_document(self.image_args(broken))
            factory.assert_not_called()

            animated = self.work / "animated.webp"
            with Image.new("RGB", (40, 40), "red") as first, Image.new("RGB", (40, 40), "blue") as second:
                first.save(animated, "WEBP", save_all=True, append_images=[second], duration=100, loop=0)
            with self.assertRaisesRegex(ValueError, "Анимированные изображения"):
                read_document.extract_document(self.image_args(animated))
            factory.assert_not_called()

            args = self.image_args(self.picture)
            args.vision_url = ""
            with self.assertRaisesRegex(ValueError, "Не подключён mmproj"):
                read_document.extract_document(args)
            args.vision_url = "http://127.0.0.1/unused"
            factory.return_value.describe.side_effect = RuntimeError("vision unavailable")
            with self.assertRaisesRegex(ValueError, "vision unavailable"):
                read_document.extract_document(args)
            factory.return_value.describe.side_effect = None
            factory.return_value.describe.return_value = ("PUMP-417", False)
            self.assertIn("PUMP-417", read_document.extract_document(args)["text"])


if __name__ == "__main__":
    unittest.main()
